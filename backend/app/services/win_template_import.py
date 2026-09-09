"""Read local Google Ads reports into reviewable template drafts."""
import csv
import re
import unicodedata
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

from fastapi import HTTPException
from pydantic import ValidationError

from app.services.win_templates import WinTemplateInput

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEMPLATES = 50


def _header(value):
    value = unicodedata.normalize("NFKD", str(value or "").lower().replace("đ", "d"))
    return re.sub(r"[^a-z0-9]+", " ", "".join(c for c in value if not unicodedata.combining(c))).strip()


def _asset_kind(header):
    match = re.fullmatch(r"(headlines?|titles?|tieu de|dau de|descriptions?|mo ta)\s*(\d*)", header)
    if not match:
        return None
    return ("descriptions" if match[1] in {"description", "descriptions", "mo ta"} else "headlines", int(match[2] or 0))


def _read_table(rows, source, filename):
    result, header, columns = [], None, []
    for number, cells in enumerate(rows, 1):
        if number > 10000:
            raise ValueError("File có quá nhiều dòng. Chỉ xuất các quảng cáo win cần nhập.")
        values = [str(cell).strip() if cell is not None else "" for cell in cells]
        if not any(values):
            continue
        if header is None:
            possible = [_header(cell) for cell in values]
            asset_columns = [(i, _asset_kind(cell)) for i, cell in enumerate(possible) if _asset_kind(cell)]
            if {kind[0] for _, kind in asset_columns} == {"headlines", "descriptions"}:
                header, columns = possible, sorted(asset_columns, key=lambda item: item[1][1])
            elif number >= 30:
                break
            continue
        def value_for(*names):
            return next((values[i] for i, name in enumerate(header) if name in names and i < len(values) and values[i]), "")
        assets = {"headlines": [], "descriptions": []}
        for index, (kind, position) in columns:
            text = values[index] if index < len(values) else ""
            if text and text not in {"--", "—"}:
                # Combined asset cells use line breaks; punctuation stays intact.
                assets[kind].extend([part.strip() for part in text.splitlines() if part.strip()])
        if not any(assets.values()):
            continue
        draft = {
            "name": value_for("name", "template name", "ten mau", "ad name", "ten quang cao") or
                    " · ".join(filter(None, [value_for("campaign", "campaign name", "chien dich"), value_for("ad group", "ad group name", "nhom quang cao")])) or
                    f"{Path(filename).stem} · {source} · dòng {number}",
            "notes": value_for("notes", "note", "ghi chu"),
            **assets,
        }
        try:
            WinTemplateInput(**draft)
            issues = []
        except ValidationError as error:
            issues = [f"{item['loc'][0]}: {item['msg'].removeprefix('Value error, ')}" for item in error.errors()]
        result.append({**draft, "source": f"{source} · dòng {number}", "issues": issues})
        if len(result) > MAX_TEMPLATES:
            raise ValueError("Tối đa 50 mẫu mỗi lần nhập. Hãy chia nhỏ file.")
    return result


def preview_win_file(content: bytes, filename: str):
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(413, "File tối đa 5 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise HTTPException(422, "Hỗ trợ Excel .xlsx hoặc .csv. File .xls cần lưu lại thành .xlsx.")
    try:
        if suffix == ".csv":
            encoding = "utf-16" if content.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            text = content.decode(encoding)
            # Choose the delimiter that exposes headline + description columns,
            # including exports with a report title/date before the header.
            result = []
            for delimiter in (",", ";", "\t"):
                rows = csv.reader(StringIO(text), delimiter=delimiter, strict=True)
                try:
                    result = _read_table(rows, "CSV", filename)
                except csv.Error:
                    continue
                if result:
                    break
        else:
            from openpyxl import load_workbook
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("File Excel quá lớn sau giải nén. Hãy xuất riêng các quảng cáo cần nhập.")
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                result = []
                for sheet in workbook.worksheets:
                    if sheet.sheet_state != "visible":
                        continue
                    # Bound reads even when the workbook advertises huge dimensions.
                    if (sheet.max_column or 0) > 100:
                        raise ValueError("File Excel tối đa 100 cột. Hãy xuất riêng các cột quảng cáo cần nhập.")
                    rows = sheet.iter_rows(max_col=min(sheet.max_column or 100, 100), values_only=True)
                    result.extend(_read_table(rows, sheet.title, filename))
                    if len(result) > MAX_TEMPLATES:
                        raise ValueError("Tối đa 50 mẫu mỗi lần nhập. Hãy chia nhỏ file.")
            finally:
                workbook.close()
        if not result:
            raise ValueError("Không tìm thấy nội dung. File cần cột Headline 1…15 và Description 1…4 (hoặc Headlines / Descriptions, mỗi nội dung một dòng trong ô).")
        # Do not treat spreadsheet formulas as ad copy or execute them.
        for row in result:
            if any(text.startswith("=") for key in ("headlines", "descriptions") for text in row[key]):
                row["issues"].append("Có ô công thức. Hãy thay công thức bằng nội dung văn bản trước khi lưu.")
        return {"drafts": result, "filename": filename}
    except HTTPException:
        raise
    except (UnicodeError, ValueError) as error:
        raise HTTPException(422, str(error) if not isinstance(error, UnicodeError) else "CSV cần mã hóa UTF-8 hoặc UTF-16.") from error
    except Exception as error:
        raise HTTPException(422, "Không đọc được file. Hãy kiểm tra định dạng và bỏ mật khẩu bảo vệ nếu có.") from error
